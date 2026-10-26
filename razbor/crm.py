"""Загрузка портфеля из выгрузки CRM (Excel .xlsx или CSV).

Названия колонок узнаются и по переменным справочника CRM (number_dog, birthdate…), и по
русским заголовкам («Номер договора», «Дата рождения»…). Лишние колонки пропускаются.
Файл читается потоком, поэтому 700 тысяч строк не занимают много памяти.
"""

from __future__ import annotations

import codecs
import csv
import io
import os
import re
from datetime import date, datetime, timedelta
from typing import Callable, Iterator

from .store import Store

# поле программы -> варианты заголовка в выгрузке
ALIASES: dict[str, list[str]] = {
    "number": ["number_dog", "номер договора", "№ договора", "договор", "номер кд"],
    "contract_date": ["date_dog", "дата договора"],
    "fio": ["fio", "фио", "фио должника", "должник", "наименование должника", "заемщик", "фио заемщика"],
    "birth": ["birthdate", "дата рождения"],
    "inn": ["inn", "инн", "инн должника"],
    "crm_id": ["id", "id договора", "id_dog"],
    "ext_id": ["id_out", "внешний id"],
    "responsible": ["operator_name", "id_user", "ответственный сотрудник", "ответственный", "оператор"],
    "status": ["status_name", "рабочий статус", "статус", "id_status"],
    "status_ov": ["status_ov_name", "статус ов"],
    "debt": ["dolg_sum", "сумма долга", "сумма долга исх", "остаток долга без гп", "остаток основного долга"],
    "debt_total": ["ocs_sum", "остаток задолженности", "остаток долга", "остаток с гп", "сумма долга с гп"],
    "principal_left": ["dolg_osn_ocs", "основной долг остаток"],
    "principal_court": ["основной долг в суд"],
    "principal_granted": ["основной долг удовлетворено", "sum_sud", "сумма по решению", "взыскано по решению",
                          "сумма по решению суда"],
    "principal_cession": ["основной долг по цессии"],
    "principal_initial": ["dolg_osn", "основной долг исх"],
    "penalties": ["peny", "пени исх"],
    "fee": ["gospohlina", "госпошлина оплаченная", "госпошлина"],
    "last_pay_date": ["date_last_payment", "дата последнего платежа"],
    "last_pay_sum": ["sum_last_payment", "сумма последнего платежа"],
    "overdue_end": ["date_stop_proch", "дата окончания просрочки по цессии", "дата окончания просрочки"],
    "creditor": ["recoverer_name", "взыскатель"],
    "court": ["sud_send", "текущий суд", "суд"],
    "suit_date": ["date_isk", "дата искового заявления"],
    "case_no": ["number_delo_il", "№ судебного дела", "номер судебного дела", "номер дела"],
    "decision_date": ["date_sud_sol", "дата решения суда"],
    "satisfied": ["satisfied", "решение суда удовлетворено"],
    "il_type": ["type_cp", "тип исполнительного документа"],
    "il_no": ["number_sa", "№ исполнительного документа", "номер исполнительного документа"],
    "il_date": ["date_sa", "дата исполнительного документа"],
    "fssp_date": ["date_fssp", "дата заявления о возбуждении ип"],
    "fssp_otdel": ["fssp_otdel_code", "отдел судебных приставов"],
    "fssp_paid": ["sum_payments_po_ip", "сумма платежей по данным фссп"],
    "fssp_actual": ["date_actuality_fssp_data", "дата актуальности данных фссп"],
    "bankrupt": ["bankrot", "банкрот"],
    "bankrupt_date": ["date_sa_bankrot", "дата решения суда о признании банкротом"],
    "bankrupt_case": ["number_bankrot_delo"],
    "death_date": ["inh_death_date", "дата смерти"],
    "estate_no": ["inh_delo_nomer", "номер наследственного дела"],
    "notary": ["inh_fio_notarius", "фио нотариуса"],
    "cession_no": ["number_act", "dr_secciyaksk", "номер акта приема-передачи", "цессия"],
    "cession_date": ["date_act", "дата акта приема-передачи"],
    "address": ["address_reg", "адрес регистрации"],
    "inn_status": ["inn_status", "статус инн"],
    "ip_status": ["dr_status_ip", "статус ип"],
    "ip_no": ["number_ip", "номер ип", "№ ип"],
    "ip_rise": ["date_rise_ip", "дата возбуждения ип"],
    "ip_renewal": ["date_renewal_ip", "дата возобновления ип"],
    "ip_end_crm": ["date_end_ip", "дата окончания ип"],
    "ip_susp": ["date_suspension_ip", "дата приостановления ип"],
    "ip_end_reason": ["end_ip_reason", "причина окончания ип", "основание окончания ип"],
}
# Колонки, которые программе не нужны и намеренно не загружаются (контакты и служебные коды)
IGNORED = {"phone", "email", "телефон", "id_operator", "id_end_ip_reason", "dr_spav_2_status_ov", "name_delo",
           "mobile", "почта", "электронная почта"}
DATE_FIELDS = {"contract_date", "birth", "last_pay_date", "overdue_end", "suit_date", "decision_date", "il_date",
               "fssp_date", "fssp_actual", "bankrupt_date", "death_date", "cession_date", "ip_rise", "ip_renewal",
               "ip_end_crm", "ip_susp"}
MONEY_FIELDS = {"debt_total", "debt", "principal_left", "principal_court", "principal_granted", "principal_cession",
                "principal_initial", "penalties", "fee", "last_pay_sum", "fssp_paid"}
RECOMMENDED = ["number", "fio", "birth", "inn", "debt", "last_pay_date", "status"]

FIELD_TITLES = {
    "number": "Номер договора", "contract_date": "Дата договора", "fio": "ФИО / наименование", "birth": "Дата рождения",
    "inn": "ИНН", "responsible": "Ответственный", "status": "Рабочий статус", "debt": "Сумма долга",
    "principal_left": "Основной долг, остаток", "principal_court": "Основной долг в суд",
    "principal_granted": "Основной долг, удовлетворено", "principal_cession": "Основной долг по цессии",
    "principal_initial": "Основной долг, исходный", "penalties": "Пени", "fee": "Госпошлина",
    "last_pay_date": "Дата последнего платежа", "last_pay_sum": "Сумма последнего платежа",
    "overdue_end": "Дата окончания просрочки", "creditor": "Взыскатель", "court": "Суд", "suit_date": "Дата иска",
    "case_no": "Номер дела", "decision_date": "Дата решения суда", "satisfied": "Результат",
    "il_type": "Тип исп. документа", "il_no": "Номер исп. документа", "il_date": "Дата исп. документа",
    "fssp_date": "Дата заявления об ИП", "fssp_otdel": "Отдел приставов", "fssp_paid": "Платежи по ИП",
    "fssp_actual": "Дата актуальности данных ФССП", "bankrupt": "Банкрот", "bankrupt_date": "Дата признания банкротом",
    "bankrupt_case": "Дело о банкротстве", "death_date": "Дата смерти", "estate_no": "Наследственное дело",
    "notary": "Нотариус", "cession_no": "Акт приёма-передачи №", "cession_date": "Акт приёма-передачи, дата",
    "address": "Адрес регистрации", "crm_id": "ID в CRM", "debt_total": "Остаток задолженности (с госпошлиной)", "ext_id": "Внешний ID", "status_ov": "Статус ОВ",
    "inn_status": "Статус ИНН", "ip_status": "Статус ИП", "ip_no": "Номер ИП", "ip_rise": "Дата возбуждения ИП",
    "ip_renewal": "Дата возобновления ИП", "ip_end_crm": "Дата окончания ИП", "ip_susp": "Дата приостановления ИП",
    "ip_end_reason": "Причина окончания ИП",
}

# ---------------------------------------------------------------- нормализация

_LATIN_TO_CYR = str.maketrans("ABCEHKMOPTXY", "АВСЕНКМОРТХУ")
_DASHES = str.maketrans({"–": "-", "—": "-", "‑": "-", "−": "-", " ": " "})


def norm_header(h) -> str:
    s = str(h or "").lower().replace("ё", "е").translate(_DASHES)
    return re.sub(r"[^0-9a-zа-я_№]+", "", s)


_IGNORED_N = {norm_header(x) for x in IGNORED}
_ALIAS_INDEX: dict[str, str] = {}
_ALIAS_PRIO: dict[str, int] = {}
for _f, _al in ALIASES.items():
    for _i, _a in enumerate(_al):
        if norm_header(_a) not in _ALIAS_INDEX:
            _ALIAS_INDEX[norm_header(_a)] = _f
            _ALIAS_PRIO[norm_header(_a)] = _i


def norm_number(s) -> str:
    """Номер договора/дела для сравнения: без пробелов, в верхнем регистре, латиница-двойники -> кириллица."""
    t = str(s or "").upper().replace("Ё", "Е").translate(_DASHES)
    t = re.sub(r"\s+", "", t)
    return t.translate(_LATIN_TO_CYR).strip(".,;")


def norm_text(s) -> str:
    t = str(s or "").lower().replace("ё", "е")
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-zа-я]+", " ", t)).strip()


_ORG = re.compile(r"(?i)\b(ооо|ао|пао|нао|оао|зао|мкк|мфк|кпк|кб|акб|пко|банк|общество|компания|кооператив)\b|[«\"]")


def detect_kind(fio: str, inn: str) -> str:
    low = (fio or "").lower().strip()
    if re.match(r"(ип|индивидуальный\s+предприниматель)\b", low):
        return "ИП"
    if _ORG.search(low):
        return "Юрлицо"
    if len(inn or "") == 10:
        return "Юрлицо"
    return "Физлицо"


def make_surname_key(fio_norm: str, kind: str) -> str:
    if kind == "Юрлицо" or not fio_norm:
        return ""
    toks = [t for t in fio_norm.split() if t not in ("ип", "индивидуальный", "предприниматель")]
    return toks[0][:5] if toks else ""


def parse_date(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        if 20000 < float(v) < 80000:
            return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
        return None
    s = str(v).strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        y, mo, d = int(m[1]), int(m[2]), int(m[3])
    else:
        m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)
        if not m:
            return None
        d, mo, y = int(m[1]), int(m[2]), int(m[3])
    try:
        if not 1900 <= y <= 2100:
            return None
        return date(y, mo, d).isoformat()
    except ValueError:
        return None


def parse_money(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace("₽", "").replace("руб.", "").replace("руб", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def clean_text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def parse_inn(v) -> str | None:
    s = re.sub(r"\D", "", str(v if v is not None else ""))
    if isinstance(v, (int, float)) and len(s) in (9, 11):
        s = s.zfill(len(s) + 1)  # Excel съел ведущий ноль
    return s if len(s) in (10, 12) else (s or None)


def parse_flag(v) -> int | None:
    s = str(v if v is not None else "").strip().lower()
    if not s:
        return None  # пусто: не затирать уже известное значение
    return 1 if s in ("1", "да", "true", "yes", "y", "банкрот", "x", "+") else 0


# ---------------------------------------------------------------- чтение файлов

csv.field_size_limit(10_000_000)


def _iter_csv(path: str) -> Iterator[list]:
    raw = open(path, "rb").read(65536)
    enc = "utf-8-sig"
    try:  # кусок файла может оборваться посреди буквы: такой обрыв не считается ошибкой
        codecs.getincrementaldecoder("utf-8")().decode(raw, final=False)
    except UnicodeDecodeError:
        enc = "cp1251"
    sample = raw.decode(enc, errors="replace")
    try:
        dialect = csv.Sniffer().sniff(sample.split("\n", 1)[0] + "\n", delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    with open(path, "r", encoding=enc, errors="replace", newline="") as f:
        for row in csv.reader(f, dialect):
            yield row


def _sheet_score(rows: list) -> int:
    return max((sum(1 for c in r if norm_header(c) in _ALIAS_INDEX) for r in rows), default=0)


def _iter_xlsx_tables(path: str) -> Iterator[tuple[str, Iterator[list]]]:
    """Все листы книги, на которых есть шапка портфеля (данные можно разбить по листам)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        good = []
        for ws in wb.worksheets:
            head = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                head.append(row)
                if i >= 12:
                    break
            if _sheet_score(head) >= 2:
                good.append(ws.title)
        if not good:
            good = [wb.worksheets[0].title]
        for title in good:
            yield title, (list(r) for r in wb[title].iter_rows(values_only=True))
    finally:
        wb.close()


def open_tables(path: str) -> Iterator[tuple[str, Iterator[list]]]:
    """(название, строки) для каждой таблицы файла: лист Excel или весь CSV."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _iter_xlsx_tables(path)
    if ext in (".csv", ".txt", ".tsv"):
        return iter([(os.path.basename(path), _iter_csv(path))])
    if ext == ".xls":
        raise ValueError("Старый формат .xls не читается. Сохраните выгрузку как .xlsx или .csv.")
    raise ValueError(f"Неизвестный формат файла: {ext}")


def open_rows(path: str) -> Iterator[list]:
    for _, rows in open_tables(path):
        return rows
    return iter([])


def map_headers(headers: list) -> tuple[dict[int, str], list[str]]:
    """Колонка -> поле. Если на одно поле подходят несколько колонок, берётся та, что выше в списке вариантов
    (например, «status_name» важнее числового «id_status»)."""
    best: dict[str, tuple[int, int]] = {}  # поле -> (приоритет, номер колонки)
    unknown: list[str] = []
    for i, h in enumerate(headers):
        if h is None or str(h).strip() == "":
            continue
        n = norm_header(h)
        if n in _IGNORED_N:
            continue
        f = _ALIAS_INDEX.get(n)
        if not f:
            unknown.append(str(h).strip())
            continue
        pr = _ALIAS_PRIO[n]
        if f not in best or pr < best[f][0]:
            if f in best:
                unknown.append(str(headers[best[f][1]]).strip())
            best[f] = (pr, i)
        else:
            unknown.append(str(h).strip())
    return {i: f for f, (_, i) in best.items()}, unknown


def build_contract(values: list, mapping: dict[int, str]) -> dict | None:
    rec: dict = {}
    for i, f in mapping.items():
        if i >= len(values):
            continue
        v = values[i]
        if isinstance(v, str) and v.strip().upper() in ("NULL", "NONE", "#N/A", "Н/Д"):
            v = None
        if f in DATE_FIELDS:
            rec[f] = parse_date(v)
        elif f in MONEY_FIELDS:
            rec[f] = parse_money(v)
        elif f == "inn":
            rec[f] = parse_inn(v)
        elif f == "bankrupt":
            rec[f] = parse_flag(v)
        else:
            rec[f] = clean_text(v)
    num = rec.get("number")
    if not num:
        return None
    if "fio" in rec or "inn" in rec:
        fio = rec.get("fio") or ""
        rec["kind"] = detect_kind(fio, rec.get("inn") or "")
    if "fio" in rec:
        rec["fio_norm"] = norm_text(rec["fio"] or "")
        rec["surname_key"] = make_surname_key(rec["fio_norm"], rec.get("kind") or "Физлицо")
    if "case_no" in rec:
        rec["case_norm"] = norm_number(rec["case_no"]) if rec.get("case_no") else None
    if "il_no" in rec:
        rec["il_norm"] = norm_number(rec["il_no"]) if rec.get("il_no") else None
    if "ip_no" in rec:
        rec["ip_norm"] = norm_number(rec["ip_no"]) if rec.get("ip_no") else None
    return rec


def import_crm(store: Store, path: str, progress: Callable[[int, str], None] | None = None,
               cancel=None, chunk: int = 5000) -> dict:
    """Загрузить выгрузку CRM. Одному договору может соответствовать несколько строк (блоки суда, ИП и т. п.):
    они сливаются, пустые ячейки не затирают уже известные значения. Листов Excel может быть несколько."""
    rows_total = added = updated = skipped = 0
    seen: set[int] = set()
    found_fields: set[str] = set()
    unknown_all: list[str] = []
    tables_used = 0
    errors: list[str] = []

    for title, it in open_tables(path):
        buf: list[list] = []
        header_idx, best = 0, -1
        for i, row in enumerate(it):
            buf.append(row)
            sc = sum(1 for c in row if norm_header(c) in _ALIAS_INDEX)
            if sc > best:
                best, header_idx = sc, i
            if i >= 14:
                break
        if best < 2:
            found = [str(c) for c in (buf[0] if buf else []) if c][:12]
            errors.append(f"«{title}»: не найдены колонки портфеля (первая строка: {', '.join(found)})")
            continue
        mapping, unknown = map_headers(buf[header_idx])
        if "number" not in mapping.values():
            errors.append(f"«{title}»: нет колонки с номером договора")
            continue
        tables_used += 1
        found_fields |= set(mapping.values())
        unknown_all += [u for u in unknown if u not in unknown_all]
        batch_rows: list[dict] = []
        batch_norm: list[str] = []
        batch_keys: list[str] = []
        use_id = "crm_id" in mapping.values()  # в CRM есть свой ID записи: он и есть ключ (номера договоров повторяются)

        def flush() -> None:
            nonlocal added, updated, batch_rows, batch_norm, batch_keys
            if batch_rows:
                a, u = store.upsert_contracts(batch_rows, batch_keys, batch_norm)
                added += a
                updated += u
                batch_rows, batch_norm, batch_keys = [], [], []

        def feed(values: list) -> None:
            nonlocal rows_total, skipped
            if not any(v not in (None, "") for v in values):
                return
            rows_total += 1
            rec = build_contract(values, mapping)
            if rec is None:
                skipped += 1
                return
            nn = norm_number(rec["number"])
            key = f"id:{rec['crm_id']}" if use_id and rec.get("crm_id") else f"n:{nn}"
            seen.add(hash(key))
            batch_rows.append(rec)
            batch_norm.append(nn)
            batch_keys.append(key)
            if len(batch_rows) >= chunk:
                flush()
                if progress:
                    progress(rows_total, "Загружаю портфель")
                if cancel is not None and cancel.is_set():
                    raise InterruptedError("Остановлено пользователем")

        for row in buf[header_idx + 1:]:
            feed(row)
        for row in it:
            feed(row)
        flush()

    if not tables_used:
        raise ValueError("В файле не найдены колонки портфеля. " + "; ".join(errors) +
                         ". Нужны хотя бы «Номер договора» и «ФИО».")
    missing = [FIELD_TITLES[f] for f in RECOMMENDED if f not in found_fields]
    store.log_import(os.path.basename(path), rows_total, added, len(seen) - added, skipped, unknown_all)
    return {"rows": rows_total, "contracts": len(seen), "added": added, "updated": len(seen) - added,
            "skipped": skipped, "tables": tables_used, "errors": errors,
            "recognized": [FIELD_TITLES[f] for f in sorted(found_fields)], "unmapped": unknown_all,
            "missing": missing}
